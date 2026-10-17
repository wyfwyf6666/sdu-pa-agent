"""教学功能路由：概念解析、论文阅读、案例分析、教案、测验、导出"""
import tempfile, os, json, io
from pathlib import Path
from fastapi import APIRouter, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional

from backend.app.services.chat import chat_service

router = APIRouter()
JSON_FMT = {"type": "json_object"}


def _extract_file_text(file: UploadFile) -> str:
    ext = os.path.splitext(file.filename or "")[1].lower()
    content_bytes = file.file.read()
    if ext == ".txt":
        return content_bytes.decode("utf-8", errors="ignore")
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        tmp.write(content_bytes)
        tmp_path = tmp.name
    try:
        if ext == ".pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(tmp_path)
                if reader.is_encrypted:
                    try: reader.decrypt("")
                    except: pass
            except ImportError:
                import PyPDF2
                reader = PyPDF2.PdfReader(tmp_path)
            return "\n".join(page.extract_text() or "" for page in reader.pages)
        elif ext in (".docx", ".doc"):
            import docx
            doc = docx.Document(tmp_path)
            return "\n".join(p.text for p in doc.paragraphs)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")
    finally:
        os.unlink(tmp_path)


# ── 概念解析（JSON输出）─────────────────────────────────

class ConceptRequest(BaseModel):
    concept: str
    context: Optional[str] = None


CONCEPT_PROMPT = """你是公共管理学科教授。对概念输出严格JSON，禁止Markdown、禁止"首先/其次/总之"等废话。

概念：{concept}{context}

输出JSON schema：
{{"definition":"2-3句定义","core_elements":[{{"element":"要素名","explanation":"说明"}}],"theory_origin":"100字理论渊源","classic_case":"中国城市治理实际案例","related_concepts":[{{"name":"相关概念","relation":"关联说明"}}],"readings":[{{"title":"文献标题","author":"作者","reason":"推荐理由"}}]}}"""


@router.post("/teaching/concept")
async def explain_concept(request: ConceptRequest):
    ctx = f"\n补充背景：{request.context}" if request.context else ""
    prompt = CONCEPT_PROMPT.format(concept=request.concept, context=ctx)
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", enable_rag=True,
            response_format=JSON_FMT)
        return {"success": True, "concept": request.concept, "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── 论文阅读助手（JSON输出）─────────────────────────────

class PaperTextRequest(BaseModel):
    text: Optional[str] = None

class PaperOutlineRequest(BaseModel):
    topic: str
    requirements: Optional[str] = None


PAPER_SUMMARY_PROMPT = """你是公共管理学科导师。对论文内容输出严格JSON速读摘要。禁止Markdown、禁止废话连词。

论文内容：
{content}

输出JSON schema：
{{"research_question":"核心研究问题","method":{{"type":"定量/定性/混合","data_source":"数据来源","sample":"样本描述"}},"findings":[{{"finding":"发现","evidence":"支撑证据"}}],"theory_contribution":"理论贡献","policy_implications":["政策启示"],"limitations":["研究局限"],"keywords":["3-5个关键词"]}}"""


PAPER_OUTLINE_PROMPT = """你是公共管理论文写作导师。输出严格JSON论文大纲。禁止Markdown、禁止废话。

选题：{topic}{requirements}

输出JSON schema：
{{"title":"学术化论文标题","abstract":"150字摘要","sections":[{{"heading":"一级标题","subsections":["子标题"],"key_points":["要点"]}}],"references_direction":[{{"title":"推荐文献","relevance":"关联说明"}}],"suggested_word_count":8000}}"""


@router.post("/teaching/paper-summary")
async def summarize_paper(request: PaperTextRequest):
    if not request.text or not request.text.strip():
        raise HTTPException(status_code=400, detail="请提供论文文本内容")
    prompt = PAPER_SUMMARY_PROMPT.format(content=request.text.strip()[:8000])
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", response_format=JSON_FMT)
        return {"success": True, "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/teaching/paper-summary-upload")
async def summarize_paper_upload(file: UploadFile = File(...)):
    try:
        content = _extract_file_text(file)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception:
        raise HTTPException(status_code=400, detail="文件解析失败")
    if not content or not content.strip():
        raise HTTPException(status_code=400, detail="未能提取到文本内容")
    prompt = PAPER_SUMMARY_PROMPT.format(content=content[:8000])
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", response_format=JSON_FMT)
        return {"success": True, "filename": file.filename,
                "extracted_length": len(content), "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/teaching/paper-outline")
async def generate_outline(request: PaperOutlineRequest):
    reqs = f"\n要求：{request.requirements}" if request.requirements else ""
    prompt = PAPER_OUTLINE_PROMPT.format(topic=request.topic, requirements=reqs)
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", enable_rag=True, response_format=JSON_FMT)
        return {"success": True, "topic": request.topic, "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── 案例分析框架（JSON输出）─────────────────────────────

class CaseAnalysisRequest(BaseModel):
    case_description: str
    analysis_dimensions: Optional[str] = None


CASE_ANALYSIS_PROMPT = """你是公共管理案例分析专家。输出严格JSON结构化分析。禁止Markdown、禁止"首先/其次/总之"等废话。

案例：
{case_description}
维度偏好：{dimensions}

输出JSON schema：
{{"background":{{"event":"事件概述","time":"时间","location":"地点","actors":["涉及主体"],"policy_context":"政策环境"}},"stakeholders":[{{"name":"利益相关方","demand":"诉求","resources":"资源/权力","strategy":"行动策略"}}],"policy_tools":[{{"tool":"工具名称","type":"规制型/激励型/信息型/组织型","assessment":"合理性评估"}}],"governance_mechanism":{{"collaboration":"多主体协作机制","information_flow":"信息流通机制","conflict_resolution":"冲突协调机制"}},"effect_evaluation":{{"goal_achievement":"目标达成度","unintended_consequences":"意外后果","sustainability":"可持续性"}},"insights":{{"theoretical":["理论启示"],"practical":["实践启示"],"scalability":"可推广性评估"}},"discussion_questions":["课堂讨论问题"]}}"""


@router.post("/teaching/case-analysis")
async def analyze_case(request: CaseAnalysisRequest):
    dims = request.analysis_dimensions or "全面分析"
    prompt = CASE_ANALYSIS_PROMPT.format(
        case_description=request.case_description[:6000], dimensions=dims)
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", enable_rag=True, response_format=JSON_FMT)
        return {"success": True, "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── 教案智能生成（JSON输出）─────────────────────────────

class TeachingPlanRequest(BaseModel):
    topic: str
    hours: Optional[int] = 2
    level: Optional[str] = "本科"
    requirements: Optional[str] = None


TEACHING_PLAN_PROMPT = """你是公共管理学科资深教授。请采用思维链方式深度生成教案。输出严格JSON，内容务必详尽充实。

课程：{topic}，{hours}学时，对象：{level}{requirements}

按以下步骤思考并生成：
1. 分析课程在公共管理学科体系中的位置
2. 确定3个知识目标、3个能力目标、2个素养目标
3. 识别2个教学重点（含教学方法与理由）和2个教学难点（含突破策略与常见误区）
4. 设计详细教学环节，每环节说明活动、意图、所需材料，时间分配合理覆盖全部课时
5. 推荐2个经典案例并说明关联性、案例简述和引导讨论问题
6. 设计2道课后作业（含详细要求与预计用时）
7. 推荐3篇延伸阅读文献（含作者、年份、核心观点）
8. 提供教学反思要点

输出JSON schema（所有数组字段不少于2项）：
{{"course_position":"学科位置分析","objectives":{{"knowledge":["知识目标1","知识目标2","知识目标3"],"ability":["能力目标1","能力目标2","能力目标3"],"literacy":["素养目标1","素养目标2"]}},"key_points":[{{"point":"重点","method":"教学方法","rationale":"为何是重点"}}],"difficult_points":[{{"point":"难点","breakthrough":"突破策略","common_misconception":"常见误区"}}],"teaching_steps":[{{"stage":"环节名","duration_min":分钟数,"activity":"详细教学活动描述","intent":"设计意图","materials":"所需材料"}}],"case_materials":[{{"name":"案例名","summary":"案例简述","relevance":"关联说明","discussion_questions":["引导问题"]}}],"homework":[{{"type":"思考题/小论文","content":"详细题目内容","requirements":"具体要求","estimated_time":"预计用时"}}],"extended_readings":[{{"title":"文献标题","author":"作者","year":"年份","key_points":"核心观点"}}],"reflection_tips":["反思提示1","反思提示2"]}}"""


@router.post("/teaching/plan")
async def generate_teaching_plan(request: TeachingPlanRequest):
    reqs = f"\n补充要求：{request.requirements}" if request.requirements else ""
    prompt = TEACHING_PLAN_PROMPT.format(
        topic=request.topic, hours=request.hours,
        level=request.level, requirements=reqs)
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", enable_rag=True, response_format=JSON_FMT)
        return {"success": True, "topic": request.topic, "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── 随堂测验生成器（JSON输出）───────────────────────────

class QuizRequest(BaseModel):
    topic: str
    difficulty: Optional[str] = "中等"
    question_types: Optional[str] = "选择题,简答题,辨析题"
    count: Optional[int] = 10
    requirements: Optional[str] = None


QUIZ_PROMPT = """你是公共管理学科命题专家。输出严格JSON，禁止Markdown和废话。

主题：{topic}，难度：{difficulty}，题型：{question_types}，共{count}题{requirements}

JSON结构：
{{"title":"测验标题","topic":"{topic}","difficulty":"{difficulty}","total_score":100,"time_minutes":分钟,"questions":[{{"id":1,"type":"choice","question":"题目","options":["A.选项A","B.选项B","C.选项C","D.选项D"],"answer":"A","scoring_points":[],"explanation":"解析"}},{{"id":2,"type":"short_answer","question":"题目","options":[],"answer":"参考答案全文","scoring_points":["得分点1","得分点2","得分点3"],"explanation":"解析"}}]}}

规则：
1. type只能是choice、short_answer、essay三者之一
2. choice题options非空、answer为单个字母、scoring_points为空数组[]
3. short_answer和essay题options为空数组[]、answer为完整文本、scoring_points为3-5个简短得分要点
4. id从1递增"""


@router.post("/teaching/quiz")
async def generate_quiz(request: QuizRequest):
    reqs = f"\n补充要求：{request.requirements}" if request.requirements else ""
    prompt = QUIZ_PROMPT.format(
        topic=request.topic, difficulty=request.difficulty,
        question_types=request.question_types, count=request.count,
        requirements=reqs)
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", enable_rag=True, response_format=JSON_FMT)
        return {"success": True, "topic": request.topic, "content": result.get("response", "")}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ── 导出功能（Word / PPT）─────────────────────────────────

class ExportRequest(BaseModel):
    json_content: str
    format: str = "docx"
    title: Optional[str] = ""


def _build_plan_docx(data: dict, title: str) -> io.BytesIO:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    doc.styles['Normal'].font.size = Pt(11)
    h = doc.add_heading(title or data.get("topic", "教案"), 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if data.get("course_position"):
        doc.add_paragraph(data["course_position"]).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_heading("一、教学目标", 1)
    obj = data.get("objectives", {})
    for key, label in [("knowledge", "知识目标"), ("ability", "能力目标"), ("literacy", "素养目标")]:
        doc.add_heading(label, 2)
        for item in obj.get(key, []):
            doc.add_paragraph(item, style='List Bullet')
    doc.add_heading("二、教学重难点", 1)
    doc.add_heading("教学重点", 2)
    for kp in data.get("key_points", []):
        doc.add_paragraph(f"{kp.get('point','')} — 方法：{kp.get('method','')}；理由：{kp.get('rationale','')}")
    doc.add_heading("教学难点", 2)
    for dp in data.get("difficult_points", []):
        doc.add_paragraph(f"{dp.get('point','')} — 突破：{dp.get('breakthrough','')}；误区：{dp.get('common_misconception','')}")
    doc.add_heading("三、教学过程", 1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Grid Accent 1'
    for i, text in enumerate(["环节", "时长(min)", "教学活动", "设计意图"]):
        table.rows[0].cells[i].text = text
    for step in data.get("teaching_steps", []):
        row = table.add_row().cells
        row[0].text = step.get("stage", "")
        row[1].text = str(step.get("duration_min", ""))
        row[2].text = step.get("activity", "")
        row[3].text = step.get("intent", "")
    doc.add_heading("四、案例素材", 1)
    for cm in data.get("case_materials", []):
        doc.add_paragraph(f"{cm.get('name','')} — {cm.get('summary','')}")
    doc.add_heading("五、课后作业", 1)
    for hw in data.get("homework", []):
        doc.add_paragraph(f"[{hw.get('type','')}] {hw.get('content','')}（{hw.get('estimated_time','')}）")
    doc.add_heading("六、推荐阅读", 1)
    for er in data.get("extended_readings", []):
        doc.add_paragraph(f"{er.get('title','')} — {er.get('author','')} ({er.get('year','')})")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _build_plan_pptx(data: dict, title: str) -> io.BytesIO:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title or data.get("topic", "教案")
    if data.get("course_position"):
        slide.placeholders[1].text = data["course_position"]
    obj = data.get("objectives", {})
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "教学目标"
    body = slide.placeholders[1].text_frame
    for key, label in [("knowledge", "知识目标"), ("ability", "能力目标"), ("literacy", "素养目标")]:
        p = body.add_paragraph(); p.text = f"\n{label}："
        for item in obj.get(key, []):
            p = body.add_paragraph(); p.text = f"• {item}"; p.level = 1
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "教学过程"
    body = slide.placeholders[1].text_frame
    for step in data.get("teaching_steps", []):
        p = body.add_paragraph()
        p.text = f"{step.get('stage','')}（{step.get('duration_min','')}min）：{step.get('activity','')}"
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "重难点与作业"
    body = slide.placeholders[1].text_frame
    for kp in data.get("key_points", []):
        p = body.add_paragraph(); p.text = f"重点：{kp.get('point','')}"
    for hw in data.get("homework", []):
        p = body.add_paragraph(); p.text = f"作业：{hw.get('content','')}"
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


@router.post("/teaching/export")
async def export_document(request: ExportRequest):
    try:
        data = json.loads(request.json_content)
    except json.JSONDecodeError:
        raise HTTPException(status_code=400, detail="JSON内容格式错误")
    fmt = request.format.lower()
    title = request.title or data.get("title") or data.get("topic") or "导出文档"
    try:
        if fmt == "docx":
            buf = _build_plan_docx(data, title)
            filename = f"{title}.docx"
            mt = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        elif fmt == "pptx":
            buf = _build_plan_pptx(data, title)
            filename = f"{title}.pptx"
            mt = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        else:
            raise HTTPException(status_code=400, detail="format必须是docx或pptx")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"文档生成失败: {str(e)}")
    from urllib.parse import quote
    return StreamingResponse(buf, media_type=mt,
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"})


# ── 简答题智能评分 ────────────────────────────────────────

class GradeShortAnswerRequest(BaseModel):
    question: str
    reference_answer: str
    user_answer: str
    scoring_points: Optional[list] = None


GRADING_PROMPT = """你是公共管理学科的评卷老师。请根据以下得分要点逐条核对学生的回答，给出评分（0-10分）。

题目：{question}
参考答案：{reference_answer}
得分要点：{scoring_points}
学生答案：{user_answer}

评分规则（按得分要点逐条核分）：
- 每条得分要点覆盖完整且准确 → 得对应分值（总分10分按要点数均分）
- 每条得分要点部分涉及但不够完整 → 得对应分值的一半
- 每条得分要点完全未涉及 → 得0分
- 如果学生答案虽与参考表述不同，但实质含义与得分要点一致，也应视为覆盖

请输出严格JSON：
{{"score": 整数0-10, "is_correct": true或false（6分及以上为true）, "feedback": "简短评语（必包含：已覆盖的得分点√、部分覆盖的△、缺失的×）", "covered_points": ["已覆盖的得分点原文"], "partial_points": ["部分覆盖的得分点原文"], "missed_points": ["未覆盖的得分点原文"]}}"""


@router.post("/teaching/grade-answer")
async def grade_short_answer(request: GradeShortAnswerRequest):
    """使用DeepSeek对简答题/辨析题进行按得分点语义评分"""
    if not request.user_answer.strip():
        return {"score": 0, "is_correct": False, "feedback": "未作答", "covered_points": [], "partial_points": [], "missed_points": []}
    sp = request.scoring_points or []
    prompt = GRADING_PROMPT.format(
        question=request.question,
        reference_answer=str(request.reference_answer),
        scoring_points=json.dumps(sp, ensure_ascii=False) if sp else "无预定义得分点，请根据参考答案自行判断关键要点",
        user_answer=request.user_answer[:2000]
    )
    try:
        result = await chat_service.get_response(
            message=prompt, llm_type="deepseek", response_format={"type": "json_object"})
        import json
        data = json.loads(result.get("response", "{}"))
        return {
            "score": int(data.get("score", 0)),
            "is_correct": data.get("is_correct", False),
            "feedback": data.get("feedback", ""),
            "covered_points": data.get("covered_points", []),
            "partial_points": data.get("partial_points", []),
            "missed_points": data.get("missed_points", []),
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
