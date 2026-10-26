"""选题建议与文献综述服务

用户输入一个选题，服务完成三件事：
1. 选题建议：LLM 评估选题的价值 / 可行性 / 范围 / 创新空间，给出细化方向；
2. 文献搜集：OpenAlex（英文）+ Serper 站内检索（中文社科：国家哲社文献中心 / 知网 / 万方 / 维普），
   返回结构化文献列表（标题 / 作者 / 年份 / 期刊 / 摘要 / DOI / 原文链接）；
3. 文献综述：基于文献元数据（标题 + 摘要）由 LLM 生成结构化综述，不读取全文，控制 token 成本。

附加能力：export_excel —— 把文献列表导出为 Excel 总结表（借鉴 paper-summary-skills 的
summary_archiving 思想：先结构化，再在结构化数据上做问答 / 存档）。
"""

from __future__ import annotations

import io
import json
import logging
import os
import re
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any, Dict, List, Optional

from openai import AsyncOpenAI

from backend.app.services.external_academic_apis import CrossrefClient, OpenAlexClient

LOGGER = logging.getLogger(__name__)

# 中文权威学术站点（Serper 站内检索目标）
CHINESE_SITES: List[tuple] = [
    ("ncpssd", "国家哲学社会科学文献中心", "site:ncpssd.org"),
    ("ncpssd_cn", "国家哲学社会科学文献中心", "site:ncpssd.cn"),
    ("cnki", "中国知网", "site:cnki.net"),
    ("wanfang", "万方数据", "site:wanfangdata.com.cn"),
    ("cqvip", "维普中文期刊", "site:cqvip.com"),
]

# 明显是服务页/首页/资讯页而非文献详情页的域名或路径片段（用于过滤中文检索噪音）
SERVICE_DOMAIN_BLACKLIST = (
    "oa.cqvip.com", "yc.cqvip.com", "super.cqvip.com",
    "cx.cnki.net", "www.cnki.net", "navi.cnki.net",
    "www.wanfangdata.com.cn/company", "journal.cqvip.com",
    "ncpssd.org/topics", "ncpssd.org/message", "ncpssd.org/journal",
)

# 搜索结果标题是站点导航/总录类时直接丢弃
TITLE_BLACKLIST_PATTERNS = (
    "期刊总录", "期刊大全", "投稿", "期刊推荐",
    "国家哲学社会科学文献中心$",
)

def _restore_openalex_abstract(inverted: Optional[Dict[str, List[int]]]) -> str:
    """OpenAlex 的摘要以倒排索引（word -> positions）存储，还原成正常文本。"""
    if not inverted:
        return ""
    position_map: Dict[int, str] = {}
    for word, positions in inverted.items():
        for pos in positions:
            position_map[pos] = word
    if not position_map:
        return ""
    return " ".join(position_map[i] for i in sorted(position_map))


def _truncate(text: str, limit: int = 300) -> str:
    text = (text or "").strip()
    return text if len(text) <= limit else text[:limit] + "…"


class TopicReviewService:
    """选题建议与文献综述核心服务。"""

    def __init__(self) -> None:
        self.client = AsyncOpenAI(
            api_key=os.getenv("OPENAI_API_KEY", ""),
            base_url=os.getenv("OPENAI_BASE_URL", ""),
        )
        self.model = os.getenv("GENERAL_MODEL_NAME", "deepseek-chat")
        self.openalex = OpenAlexClient()
        self.serper_key = os.getenv("SERPER_API_KEY", "")

    # =========================================================================
    # 1. 选题建议
    # =========================================================================
    async def suggest_topic(self, topic: str) -> str:
        """LLM 给出选题建议（可行性 / 价值 / 范围 / 创新 / 细化方向）。"""
        system_prompt = (
            "你是一位公共管理学科的研究生导师和学术顾问，擅长帮助研究者打磨选题。"
            "请对用户给出的选题给出专业、具体、可操作的建议。"
            "要求：用中文回答，使用 Markdown 格式，控制在 500 字以内。"
        )
        user_prompt = (
            f"我准备研究以下选题，请从【选题价值】【可行性】【范围界定】"
            f"【创新空间】【可能的细化方向】【建议优先关注的文献领域】六个方面给出建议：\n\n{topic}"
        )
        content = await self._chat(system_prompt, user_prompt, temperature=0.7, max_tokens=2000)
        return content

    # =========================================================================
    # 2. 文献搜集（英文 OpenAlex + 中文 Serper）
    # =========================================================================
    async def search_literature(
        self,
        topic: str,
        en_limit: int = 8,
        zh_limit: int = 8,
    ) -> List[Dict[str, Any]]:
        """按选题搜集文献，返回统一结构的文献列表。

        流程：先由 LLM 提取中英文搜索关键词（短词），英文用关键词搜 OpenAlex，
        中文用关键词对权威中文学术站点做 Serper 站内检索，并过滤服务页噪音。
        """
        papers: List[Dict[str, Any]] = []
        topic_clean = topic.strip()

        # 2.0 提取搜索关键词（中英文各一组）
        keywords = await self._extract_keywords(topic_clean)
        en_keywords = keywords.get("en") or [topic_clean]
        zh_keywords = keywords.get("zh") or [topic_clean]
        LOGGER.info("topic review keywords: en=%s zh=%s", en_keywords, zh_keywords)

        # 2.1 英文文献（OpenAlex，按关键词逐个检索）
        try:
            en_seen: set = set()
            for kw in en_keywords[:3]:
                en_results = await asyncio_to_thread(self.openalex.search_works, kw)
                for item in (en_results or []):
                    raw = item.get("raw") or {}
                    abstract = _restore_openalex_abstract(raw.get("abstract_inverted_index"))
                    primary = raw.get("primary_location") or {}
                    oa = raw.get("open_access") or {}
                    url = (
                        primary.get("landing_page_url")
                        or oa.get("oa_url")
                        or item.get("url")
                        or ""
                    )
                    doi = item.get("doi") or ""
                    if doi and (not url or url == item.get("url")):
                        url = f"https://doi.org/{doi}"
                    title = item.get("title", "") or ""
                    key = re.sub(r"[\s\W_]+", "", title.lower())
                    if not key or key in en_seen:
                        continue
                    en_seen.add(key)
                    papers.append({
                        "title": title,
                        "authors": item.get("authors") or [],
                        "year": item.get("year"),
                        "journal": item.get("journal", ""),
                        "abstract": _truncate(abstract, 400),
                        "doi": doi,
                        "url": url,
                        "source": "OpenAlex",
                        "language": "en",
                    })
                    if sum(1 for p in papers if p.get("language") == "en") >= en_limit:
                        break
                if sum(1 for p in papers if p.get("language") == "en") >= en_limit:
                    break
        except Exception as exc:  # 英文搜索失败不阻塞整体流程
            LOGGER.warning("OpenAlex keyword search failed: %s", exc)

        # 2.2 中文文献（Serper 站内检索，多关键词合并）
        if self.serper_key:
            try:
                zh_results = await asyncio_to_thread(
                    self._serper_keyword_search, zh_keywords, zh_limit
                )
                papers.extend(zh_results)
            except Exception as exc:
                LOGGER.warning("Serper Chinese search failed: %s", exc)

        # 2.2.1 中文文献作者补全（Serper 结果不带作者字段，用 LLM 从摘要片段提取）
        zh_no_author = [p for p in papers if p.get("language") == "zh" and not p.get("authors")]
        if zh_no_author:
            try:
                await self._fill_zh_authors(zh_no_author)
            except Exception as exc:
                LOGGER.warning("Chinese authors fill failed: %s", exc)

        # 去重（按标题归一化）
        seen: set = set()
        deduped: List[Dict[str, Any]] = []
        for p in papers:
            key = re.sub(r"[\s\W_]+", "", (p.get("title") or "").lower())
            if not key or key in seen:
                continue
            seen.add(key)
            deduped.append(p)

        # 英文 / 中文分别限数
        en_papers = [p for p in deduped if p.get("language") == "en"][:en_limit]
        zh_papers = [p for p in deduped if p.get("language") == "zh"][:zh_limit]
        return en_papers + zh_papers

    async def _extract_keywords(self, topic: str) -> Dict[str, List[str]]:
        """让 LLM 从选题中提取中英文搜索关键词（各 2-3 个短词），用于提升检索命中率。"""
        try:
            system_prompt = (
                "你是文献检索专家。根据用户选题提取检索关键词，返回严格 JSON："
                '{"en": ["英文关键词1", "英文关键词2"], "zh": ["中文关键词1", "中文关键词2"]}。\n'
                "要求：\n"
                "1. en 字段必须输出英文单词或短语（例如 policy diffusion、grassroots governance），"
                "即使选题是纯中文，也要给出对应的常用英文学术术语；严禁在 en 字段输出中文字符。\n"
                "2. zh 字段输出中文关键词。\n"
                "3. 每个关键词 2-5 个词，覆盖选题核心概念；不要输出 JSON 以外的任何内容。"
            )
            content = await self._chat(system_prompt, topic, temperature=0.2, max_tokens=200)
            content = content.strip()
            # 容错：去除可能的 markdown 代码块标记
            content = re.sub(r"^```(?:json)?|```$", "", content, flags=re.MULTILINE).strip()
            data = json.loads(content)
            cjk_re = re.compile(r"[\u4e00-\u9fff]")
            en_list = [str(k).strip() for k in (data.get("en") or []) if str(k).strip()]
            # 防御：en 字段混入中文 → 剔除
            en_list = [k for k in en_list if not cjk_re.search(k)][:3]
            return {
                "en": en_list,
                "zh": [str(k).strip() for k in (data.get("zh") or []) if str(k).strip()][:3],
            }
        except Exception as exc:
            LOGGER.warning("keyword extraction failed, fallback to raw topic: %s", exc)
            return {"en": [topic[:100]], "zh": [topic[:60]]}

    def _serper_keyword_search(
        self, keywords: List[str], limit: int = 8
    ) -> List[Dict[str, Any]]:
        """通过 Serper 在权威中文学术库站内做关键词检索，返回文献页链接并过滤服务页噪音。

        - 用 LLM 提取的短关键词（而非整句选题）搜索，命中率更高；
        - 过滤已知服务页/首页域名（查重、投稿、推广等），只保留文献详情页类结果。
        """
        import http.client

        keywords = [k for k in keywords if k][:3]
        if not keywords:
            return []

        def _is_paper_page(link: str, provider_id: str) -> bool:
            """按来源站点判断链接是否为「具体文献详情页」，而非期刊门户/排行/导航页。

            - 知网：仅保留 kns.cnki.net 详情页 或 cbpt.cnki.net 下含 /paper/ 的论文详情页；
            - 万方：仅保留 d.wanfangdata.com.cn/periodical/ 论文详情页；
            - 维普：仅保留含 Article 的论文详情页；
            - 国家哲社：根域下除资讯/期刊导航页（黑名单已过滤）外的页面。
            """
            low = (link or "").lower()
            if any(b in low for b in SERVICE_DOMAIN_BLACKLIST):
                return False
            if provider_id == "cnki":
                return "kns.cnki.net" in low or "/paper/" in low
            if provider_id == "wanfang":
                return "/periodical/" in low
            if provider_id == "cqvip":
                return "article" in low
            return True

        def _one_site(site_tuple: tuple, kw: str):
            provider_id, provider_name, site_prefix = site_tuple
            out: List[Dict[str, Any]] = []
            query = f'{site_prefix} "{kw}"'
            try:
                conn = http.client.HTTPSConnection("google.serper.dev", timeout=8)
                payload = json.dumps({"q": query[:300], "gl": "cn", "hl": "zh-cn"})
                headers = {
                    "X-API-KEY": self.serper_key,
                    "Content-Type": "application/json",
                }
                conn.request("POST", "/search", payload, headers)
                resp = conn.getresponse()
                if resp.status == 200:
                    data = json.loads(resp.read().decode("utf-8"))
                    for item in data.get("organic", [])[:3]:
                        link = item.get("link", "") or ""
                        title = item.get("title", "") or ""
                        if not link or not _is_paper_page(link, provider_id):
                            continue
                        for pattern in TITLE_BLACKLIST_PATTERNS:
                            if re.search(pattern, title or ""):
                                link = ""  # 标题是站点导航类 → 丢弃
                                break
                        if not link:
                            continue
                        out.append({
                            "title": title,
                            "authors": [],
                            "year": None,
                            "journal": provider_name,
                            "abstract": _truncate(item.get("snippet", ""), 300),
                            "doi": "",
                            "url": link,
                            "source": provider_id,
                            "language": "zh",
                        })
                conn.close()
            except Exception as exc:
                LOGGER.debug("Serper keyword search %s (%s) failed: %s", provider_id, kw, exc)
            return out

        candidates: List[Dict[str, Any]] = []
        with ThreadPoolExecutor(max_workers=len(CHINESE_SITES)) as pool:
            futures = [
                pool.submit(_one_site, s, kw)
                for s in CHINESE_SITES
                for kw in keywords
            ]
            for fut in as_completed(futures, timeout=15):
                try:
                    candidates.extend(fut.result(timeout=1))
                except Exception:
                    pass
        return candidates[:limit]

    @staticmethod
    def _clean_zh_author_list(authors: List[str]) -> List[str]:
        """清洗作者列表：按逗号/顿号拆分，去除中文姓名间的 ASCII 空格，统一返回顿号分隔的姓名列表。"""
        out: List[str] = []
        for a in authors or []:
            for part in re.split(r"[、,，]", a):
                part = re.sub(r"(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])", "", part).strip()
                if part:
                    out.append(part)
        return out[:3]

    async def _fill_zh_authors(self, papers: List[Dict[str, Any]]) -> None:
        """补齐中文文献作者，三级策略：

        1. 摘要含 DOI → Crossref DOI 反查（最准确）；
        2. 无 DOI → 按标题反查 Crossref；
        3. 仍未命中 → LLM 从摘要提取 + 结合模型知识补知名文献作者。
        """
        if not papers:
            return

        # 1) DOI 反查 Crossref
        doi_re = re.compile(r"10\.\d{4,9}/[^\s,;）)\]\"'，。]+")
        crossref = CrossrefClient()
        unresolved: List[Dict[str, Any]] = []
        for p in papers:
            if p.get("authors"):
                continue
            m = doi_re.search(p.get("abstract", "") or "")
            if m:
                try:
                    result = await asyncio_to_thread(
                        crossref.search_by_doi, m.group(0).rstrip(".")
                    )
                    if result and result.get("found") and result.get("authors"):
                        p["authors"] = self._clean_zh_author_list(result["authors"])
                        continue
                except Exception as exc:
                    LOGGER.debug("Crossref DOI lookup failed: %s", exc)
            unresolved.append(p)

        # 2) 无 DOI 或反查失败 → 按标题反查 Crossref
        still: List[Dict[str, Any]] = []
        for p in unresolved:
            if p.get("authors"):
                continue
            title = (p.get("title") or "").strip()
            if not title:
                continue
            try:
                result = await asyncio_to_thread(
                    crossref.search_by_title_author_year, title
                )
                if result and result.get("found") and result.get("authors"):
                    p["authors"] = self._clean_zh_author_list(result["authors"])
                    continue
            except Exception as exc:
                LOGGER.debug("Crossref title lookup failed: %s", exc)
            still.append(p)

        # 3) 仍未补上 → LLM 从摘要提取（可结合模型知识，不确定必须留空）
        if not still:
            return
        items = "\n".join(
            f"{i}. 标题：{p.get('title', '')}\n摘要：{_truncate(p.get('abstract', ''), 200)}"
            for i, p in enumerate(still, 1)
        )
        system_prompt = (
            "你是文献信息标注助手。根据编号列表中每条文献的标题和摘要片段，判断其作者。\n"
            "提取规则：\n"
            "1. 优先从摘要中明显出现的作者信息提取（如「作者姓名. 标题[J]」「作者姓名：」"
            "「作者姓名[作者简介]」等引用或署名格式）；\n"
            "2. 若摘要中没有作者信息，可结合你自身知识判断该文献的真实作者——"
            "只有当你非常确定时才填入（如知名学者、教科书级文献），不确定时该条输出空字符串，严禁编造；\n"
            "3. 多个作者用顿号分隔，最多取 3 个；\n"
            "4. 只返回严格 JSON 对象：{\"1\": \"第一条的作者\", \"2\": \"第二条的作者\"}，"
            "key 使用输入编号，不要输出 JSON 以外的任何内容。"
        )
        content = await self._chat(system_prompt, items, temperature=0.2, max_tokens=800)
        content = re.sub(r"^```(?:json)?|```$", "", content.strip(), flags=re.MULTILINE).strip()
        try:
            mapping = json.loads(content)
        except Exception as exc:
            LOGGER.warning("authors mapping parse failed: %s", exc)
            return
        if not isinstance(mapping, dict):
            return
        filled = 0
        for i, p in enumerate(still, 1):
            author = str(mapping.get(str(i)) or "").strip()
            if author:
                p["authors"] = self._clean_zh_author_list([author])
                filled += 1
        LOGGER.info("zh authors filled %d/%d via LLM", filled, len(still))

    # =========================================================================
    # 3. 文献综述生成
    # =========================================================================
    async def generate_review(self, topic: str, papers: List[Dict[str, Any]]) -> str:
        """基于文献元数据（标题 + 摘要）生成结构化文献综述。"""
        if not papers:
            return "未能检索到足够的文献，请尝试更换关键词或扩大选题范围后再试。"

        lines = [f"用户选题：{topic}\n", "已检索到的文献（含标题/作者/年份/期刊/摘要）：\n"]
        for i, p in enumerate(papers, 1):
            authors = "、".join((p.get("authors") or [])[:3]) or "佚名"
            lines.append(
                f"[{i}] {p.get('title', '')}\n"
                f"作者：{authors}；年份：{p.get('year') or '未知'}；"
                f"来源：{p.get('journal', '')}（{p.get('source', '')}）\n"
                f"摘要：{_truncate(p.get('abstract', ''), 300)}\n"
            )
        ref_text = "\n".join(lines)

        system_prompt = (
            "你是一位公共管理学科的学术综述专家。请基于用户选题和检索到的文献元数据，"
            "撰写一篇结构清晰、内容完整的文献综述。要求：\n"
            "1. 用中文回答，使用 Markdown 格式；\n"
            "2. 结构包含：研究概况、主题脉络（按研究主题分组梳理）、核心观点对比、"
            "主要研究方法、现有研究不足、对本选题的建议，最后以一段简短的总结收尾；\n"
            "3. 只基于提供的文献信息进行归纳，不编造不存在的文献或观点；\n"
            "4. 若文献数量有限，请明确说明综述的覆盖范围有限；\n"
            "5. 务必保证文本完整，所有章节都要写到，并以完整的总结段落结束，不要中途截断。"
        )
        content = await self._chat(system_prompt, ref_text, temperature=0.3, max_tokens=4000)
        return content

    # =========================================================================
    # 4. Excel 文献总结表导出
    # =========================================================================
    def export_excel(self, papers: List[Dict[str, Any]]) -> bytes:
        """把文献列表导出为 Excel 总结表（每篇一行，含原文链接）。"""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        wb = Workbook()
        ws = wb.active
        ws.title = "文献总结表"

        headers = ["序号", "标题", "作者", "年份", "期刊/来源", "摘要", "DOI", "原文链接"]
        ws.append(headers)
        for col, _ in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(vertical="center")

        for i, p in enumerate(papers, 1):
            ws.append([
                i,
                p.get("title", ""),
                "、".join((p.get("authors") or [])[:5]),
                p.get("year") or "",
                p.get("journal", ""),
                p.get("abstract", ""),
                p.get("doi", ""),
                p.get("url", ""),
            ])

        widths = [6, 50, 30, 8, 22, 60, 28, 50]
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()

    # =========================================================================
    # 通用 LLM 调用
    # =========================================================================
    async def _chat(
        self,
        system_prompt: str,
        user_prompt: str,
        temperature: float = 0.3,
        max_tokens: int = 1500,
    ) -> str:
        resp = await self.client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=temperature,
            max_tokens=max_tokens,
        )
        return (resp.choices[0].message.content or "").strip()


def asyncio_to_thread(fn, *args, **kwargs):
    """同步函数丢到线程池执行，避免阻塞 FastAPI 事件循环。"""
    import asyncio

    return asyncio.to_thread(fn, *args, **kwargs)


topic_review_service = TopicReviewService()
